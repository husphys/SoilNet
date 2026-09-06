#!/usr/bin/env python3
"""Schema-only audit of local historical CSV/XLSX results; copies no values."""
from __future__ import annotations

import csv
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
from soilnet.io import resolve_paths, write_csv


def main() -> int:
    root = resolve_paths()["data_root"]
    rows = []
    for path in sorted((p for p in root.rglob("*") if p.is_file() and p.suffix.casefold() in {".csv", ".xlsx"}), key=lambda p: p.as_posix().casefold()):
        columns, row_count, error = [], "", ""
        try:
            if path.suffix.casefold() == ".csv":
                with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
                    reader = csv.reader(handle)
                    columns = next(reader, [])
                    row_count = sum(1 for _ in reader)
            else:
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True)
                sheet = workbook[workbook.sheetnames[0]]
                iterator = sheet.iter_rows(values_only=True)
                columns = [str(value) if value is not None else "" for value in next(iterator, [])]
                row_count = sum(1 for _ in iterator)
                workbook.close()
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
        joined = ";".join(columns)
        lower = joined.casefold()
        scope = "TRAINING_LOG_SUSPECTED" if any(token in lower for token in ("epoch", "train_loss", "training")) else "SCOPE_NOT_PROVEN"
        rows.append({
            "relative_path": path.relative_to(root).as_posix(),
            "file_type": path.suffix.casefold().lstrip("."),
            "row_count": row_count,
            "columns": joined,
            "evaluation_scope": scope,
            "read_error": error,
        })
    write_csv(REPO / "results" / "audit" / "historical_result_schema.csv", rows, [
        "relative_path", "file_type", "row_count", "columns", "evaluation_scope", "read_error",
    ])
    print(f"Audited schemas for {len(rows)} local CSV/XLSX files")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
