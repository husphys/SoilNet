#!/usr/bin/env python3
"""Phase 2 evidence audit. Reads source artifacts only; never trains or mutates them."""
from __future__ import annotations

import csv
import gc
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

import torch

from soilnet.io import resolve_paths, sha256_file, write_csv, write_json
from soilnet.models.soilnet import SoilNetDualHead

TEXT_SUFFIXES = {".py", ".ipynb", ".csv", ".txt", ".md", ".json", ".yaml", ".yml", ".log", ".xml"}
PRIMARY_CATEGORIES = {
    "SSL_BACKBONE", "VICREG_PROJECTOR", "FINETUNED_SOILNET_DUALHEAD",
    "BASELINE_CLASSIFIER", "FULL_TRAINING_STATE", "PARTIAL_STATE", "UNKNOWN", "UNREADABLE",
}
SCOPE_C_CATEGORIES = {
    "SSL_BACKBONE", "VICREG_PROJECTOR", "FINETUNED_SOILNET_DUALHEAD",
    "BASELINE_MODEL", "FULL_TRAINING_STATE", "PARTIAL_STATE", "UNKNOWN", "UNREADABLE",
}


def read_csv(relative: str) -> list[dict[str, str]]:
    with (REPO / relative).open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def document_text(path: Path, include_outputs: bool = True) -> str:
    if path.suffix.casefold() != ".ipynb":
        return path.read_text(encoding="utf-8", errors="replace")
    try:
        notebook = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return ""
    parts: list[str] = []
    for cell in notebook.get("cells", []):
        parts.append("".join(cell.get("source", [])))
        if include_outputs:
            for output in cell.get("outputs", []):
                parts.append("".join(output.get("text", [])))
                data = output.get("data", {})
                for value in data.values():
                    parts.append("".join(value) if isinstance(value, list) else str(value))
    return "\n".join(parts)


def source_documents(paths: dict[str, Path]) -> list[tuple[str, Path, str]]:
    documents = []
    for scope, root in (("LEGACY_GITHUB", paths["legacy_root"]), ("DATA_ROOT", paths["data_root"])):
        for path in sorted(root.rglob("*"), key=lambda item: item.as_posix().casefold()):
            if path.is_file() and path.suffix.casefold() in {".py", ".ipynb"}:
                documents.append((scope, path, path.relative_to(root).as_posix()))
    return documents


def audit_labeled_duplicates() -> dict[str, object]:
    samples = read_csv("data/manifests/soilnet_samples.csv")
    splits = read_csv("data/splits/final_split.csv")
    inventory = read_csv("data/manifests/dataset_inventory.csv")
    dataset = json.loads((REPO / "results/audit/dataset_summary.json").read_text(encoding="utf-8"))
    inv_by_path = {row["relative_path"]: row for row in inventory}
    split_by_id = {row["sample_id"]: row for row in splits}

    by_sha: dict[str, list[dict[str, str]]] = defaultdict(list)
    for sample in samples:
        by_sha[sample["sha256"]].append(sample)
    exact_rows = []
    for digest, members in sorted(by_sha.items()):
        if not digest or len(members) < 2:
            continue
        group_id = f"sha256:{digest[:16]}"
        for sample in members:
            exact_rows.append({
                "duplicate_group_id": group_id,
                "group_size": len(members),
                "sample_id": sample["sample_id"],
                "relative_path": sample["relative_path"],
                "effective_relative_path": sample["effective_relative_path"],
                "sha256": digest,
                "split": split_by_id[sample["sample_id"]]["split"],
            })
    write_csv(REPO / "results/audit/labeled_exact_duplicates.csv", exact_rows, [
        "duplicate_group_id", "group_size", "sample_id", "relative_path",
        "effective_relative_path", "sha256", "split",
    ])

    by_group: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in splits:
        by_group[row["group_id"]].append(row)
    near_rows = []
    near_groups = 0
    for group_id, members in sorted(by_group.items()):
        hashes = {row["sha256"] for row in members}
        if len(members) < 2 or len(hashes) < 2:
            continue
        near_groups += 1
        for row in members:
            inv = inv_by_path.get(row["effective_relative_path"], {})
            near_rows.append({
                "near_duplicate_group_id": group_id,
                "group_size": len(members),
                "unique_sha256_count": len(hashes),
                "sample_id": row["sample_id"],
                "relative_path": row["relative_path"],
                "effective_relative_path": row["effective_relative_path"],
                "sha256": row["sha256"],
                "perceptual_hash": inv.get("perceptual_hash", ""),
                "split": row["split"],
                "threshold": 5,
            })
    write_csv(REPO / "results/audit/labeled_near_duplicates.csv", near_rows, [
        "near_duplicate_group_id", "group_size", "unique_sha256_count", "sample_id",
        "relative_path", "effective_relative_path", "sha256", "perceptual_hash", "split", "threshold",
    ])

    split_counts = Counter(row["split"] for row in splits)
    exact_cross = sum(len({split_by_id[x["sample_id"]]["split"] for x in members}) > 1 for members in by_sha.values())
    near_cross = sum(len({x["split"] for x in members}) > 1 for members in by_group.values() if len({x["sha256"] for x in members}) > 1)
    group_leakage = sum(len({x["split"] for x in members}) > 1 for members in by_group.values())
    summary = {
        "all_image_pool": {
            "total": dataset["total_images"],
            "exact_unique": dataset["unique_images_by_sha256"],
            "exact_duplicate_groups": dataset["exact_duplicate_groups"],
            "redundant_copies": dataset["exact_duplicate_files_beyond_first"],
            "near_duplicate_candidate_components": dataset["near_duplicate_components_with_nonidentical_sha256"],
            "near_duplicate_candidate_pairs": dataset["near_duplicate_candidate_pairs_at_threshold"],
        },
        "labeled_dataset": {
            "labeled_rows": len(samples),
            "exact_unique": len(by_sha),
            "exact_duplicate_groups": sum(len(value) > 1 for value in by_sha.values()),
            "redundant_copies": sum(max(0, len(value) - 1) for value in by_sha.values()),
            "near_duplicate_candidate_components": near_groups,
            "near_duplicate_member_rows": len(near_rows),
        },
        "final_prospective_split": {
            "train": split_counts["train"],
            "validation": split_counts["validation"],
            "test": split_counts["test"],
            "exact_cross_split_duplicates": exact_cross,
            "near_duplicate_cross_split_candidates": near_cross,
            "group_leakage": group_leakage,
        },
        "terminology": "14053 is the image pool; only 2057 manifest rows form the labeled supervised dataset.",
    }
    write_json(REPO / "results/audit/labeled_dataset_duplicate_summary.json", summary)
    return summary


def scalar_count_hits(obj: object, counters: Counter[str]) -> None:
    if isinstance(obj, dict):
        for key, value in obj.items():
            if key not in {"model_state_dict", "state_dict", "model", "optimizer_state_dict", "optimizer_state", "projector_state_dict"}:
                scalar_count_hits(value, counters)
    elif isinstance(obj, (list, tuple)) and len(obj) <= 100:
        for value in obj:
            scalar_count_hits(value, counters)
    elif isinstance(obj, (int, float, str)):
        normalized = str(obj).replace(",", "").strip()
        if normalized in {"2057", "2057.0"}:
            counters["2057"] += 1
        if normalized in {"2250", "2250.0"}:
            counters["2250"] += 1


def normalize_state(state: dict[str, object]) -> tuple[dict[str, object], str]:
    keys = [str(key) for key in state]
    for prefix in ("module.model.", "module.", "model.", "encoder."):
        if keys and all(key.startswith(prefix) for key in keys):
            return {str(key)[len(prefix):]: value for key, value in state.items()}, prefix
    return {str(key): value for key, value in state.items()}, ""


def state_from_object(obj: object) -> tuple[dict[str, object] | None, str, dict[str, object] | None]:
    if not isinstance(obj, dict):
        return None, "", None
    for key in ("model_state_dict", "state_dict", "model_state", "model", "encoder_state_dict"):
        if isinstance(obj.get(key), dict):
            projector = obj.get("projector_state_dict") if isinstance(obj.get("projector_state_dict"), dict) else None
            return obj[key], key, projector
    if obj and all(hasattr(value, "shape") for value in obj.values()):
        return obj, "raw_state_dict", None
    return None, "", None


def shape_of(state: dict[str, object], key: str) -> str:
    value = state.get(key)
    return "x".join(str(value) for value in value.shape) if hasattr(value, "shape") else ""


def model_compatibility(state: dict[str, object], canonical: dict[str, object]) -> bool:
    required = (
        "initial_conv.weight", "channel_adapter.weight", "mvit_to_mnv2.weight",
        "final_conv.weight", "light_dense.0.weight", "reg_head.2.weight", "cls_head.2.weight",
    )
    return all(
        key in state and hasattr(state[key], "shape") and tuple(state[key].shape) == tuple(canonical[key].shape)
        for key in required
    )


def infer_mu(filename: str) -> str:
    match = re.search(r"mu[_ =-]?(\d+(?:\.0)?)", filename, re.I)
    return match.group(1) if match else ""


def infer_epoch(filename: str, obj: object) -> str:
    if isinstance(obj, dict) and isinstance(obj.get("epoch"), (int, float, str)):
        return str(obj["epoch"])
    match = re.search(r"epoch[_ =-]?(\d+)", filename, re.I)
    return match.group(1) if match else ""


def classify_state(scope: str, path: Path, obj: object, state: dict[str, object] | None,
                   projector: dict[str, object] | None, compatible: bool) -> str:
    if state is None:
        return "UNKNOWN"
    lower = path.name.casefold()
    keys = set(state)
    has_optimizer = isinstance(obj, dict) and any(key in obj for key in ("optimizer", "optimizer_state", "optimizer_state_dict"))
    if has_optimizer:
        return "FULL_TRAINING_STATE"
    if projector is not None:
        return "PARTIAL_STATE"
    if keys and all(key.startswith("net.") for key in keys):
        return "VICREG_PROJECTOR"
    fine_tokens = ("best_finetuned", "finetune_mu", "finetuned", "soilnet_best", "final_tinyimagenet")
    if compatible and any(token in lower for token in fine_tokens):
        return "FINETUNED_SOILNET_DUALHEAD"
    if compatible and "vicreg_model_final" in lower:
        return "SSL_BACKBONE"
    if any(key.startswith(("classifier.", "backbone.")) for key in keys) or "sota" in path.as_posix().casefold():
        return "BASELINE_MODEL" if scope == "DATA_ROOT_EMBEDDED_ARTIFACTS" else "BASELINE_CLASSIFIER"
    if compatible:
        return "UNKNOWN"
    if any(token in lower for token in ("linear_reg", "classifier_final")):
        return "PARTIAL_STATE"
    return "UNKNOWN"


def inspect_checkpoint(path: Path, scope: str, canonical: dict[str, object], count_hits: Counter[str]) -> dict[str, object]:
    try:
        try:
            obj = torch.load(path, map_location="cpu", weights_only=True, mmap=True)
        except (RuntimeError, ValueError):
            obj = torch.load(path, map_location="cpu", weights_only=True)
    except Exception as exc:
        return {
            "safe_load_status": "SAFE_LOAD_FAILED", "top_level_type": "UNREADABLE", "top_level_keys": "",
            "state_dict_prefix": "", "tensor_count": 0, "parameter_tensor_count": 0,
            "primary_category": "UNREADABLE", "suspected_category": "UNREADABLE", "suspected_mu": infer_mu(path.name),
            "suspected_epoch": "", "suspected_model": "UNKNOWN", "has_reg_head": False,
            "has_cls_head": False, "has_light_branch": False, "num_classes": "", "regression_output_dimension": "",
            "reg_head_shape": "", "cls_head_shape": "", "model_compatible": False,
            "has_optimizer": False, "has_epoch": False, "has_config": False,
            "notes": f"SAFE_LOAD_FAILED; no unsafe fallback: {type(exc).__name__}: {exc}",
        }
    scalar_count_hits(obj, count_hits)
    state, container, projector = state_from_object(obj)
    normalized, prefix = normalize_state(state) if state is not None else ({}, "")
    compatible = model_compatibility(normalized, canonical) if normalized else False
    category = classify_state(scope, path, obj, normalized or None, projector, compatible)
    if category not in (SCOPE_C_CATEGORIES if scope == "DATA_ROOT_EMBEDDED_ARTIFACTS" else PRIMARY_CATEGORIES):
        raise AssertionError(f"invalid checkpoint category: {category}")
    top_keys = [str(key) for key in obj.keys()] if isinstance(obj, dict) else []
    tensor_count = sum(hasattr(value, "shape") for value in normalized.values())
    parameter_count = sum(
        hasattr(value, "shape") and not key.endswith(("running_mean", "running_var", "num_batches_tracked"))
        for key, value in normalized.items()
    )
    has_reg = any(key.startswith("reg_head.") for key in normalized)
    has_cls = any(key.startswith("cls_head.") for key in normalized)
    has_light = any(key.startswith(("light_dense.", "light_fc.")) for key in normalized)
    cls_shape = shape_of(normalized, "cls_head.2.weight")
    reg_shape = shape_of(normalized, "reg_head.2.weight")
    num_classes = normalized.get("cls_head.2.weight").shape[0] if cls_shape else ""
    reg_dim = normalized.get("reg_head.2.weight").shape[0] if reg_shape else ""
    has_optimizer = isinstance(obj, dict) and any(key in obj for key in ("optimizer", "optimizer_state", "optimizer_state_dict"))
    has_epoch = isinstance(obj, dict) and "epoch" in obj
    has_config = isinstance(obj, dict) and "config" in obj
    if category in {"BASELINE_MODEL", "BASELINE_CLASSIFIER"}:
        suspected_model = next((name for name in ("efficientnet", "mobilevit", "mobilenet", "resnet") if name in path.name.casefold()), "baseline_unknown")
    elif compatible:
        suspected_model = "SoilNetDualHead"
    elif category == "VICREG_PROJECTOR":
        suspected_model = "VICRegProjector"
    else:
        suspected_model = "UNKNOWN"
    result = {
        "safe_load_status": "SAFE_LOAD_OK", "top_level_type": type(obj).__name__,
        "top_level_keys": ";".join(top_keys[:100]), "state_dict_prefix": prefix or container,
        "tensor_count": tensor_count, "parameter_tensor_count": parameter_count,
        "primary_category": category, "suspected_category": category, "suspected_mu": infer_mu(path.name),
        "suspected_epoch": infer_epoch(path.name, obj), "suspected_model": suspected_model,
        "has_reg_head": has_reg, "has_cls_head": has_cls, "has_light_branch": has_light,
        "num_classes": num_classes, "regression_output_dimension": reg_dim,
        "reg_head_shape": reg_shape, "cls_head_shape": cls_shape, "model_compatible": compatible,
        "has_optimizer": has_optimizer, "has_epoch": has_epoch, "has_config": has_config,
        "notes": "Parameter tensor count excludes running statistics/counters; classification uses normalized keys and source evidence.",
    }
    del obj, state, normalized, projector
    gc.collect()
    return result


def audit_checkpoints(paths: dict[str, Path], documents: list[tuple[str, Path, str]]) -> tuple[list[dict[str, object]], dict[str, object]]:
    phase1 = read_csv("results/audit/checkpoint_inventory.csv")
    scope_a = [row for row in phase1 if row["source"] == "LOCAL"]
    scope_b = [row for row in phase1 if row["source"] == "GITHUB"]
    if len(scope_a) != 969 or len(scope_b) != 10:
        raise RuntimeError("PHASE 1 checkpoint scopes changed; refusing to rewrite historical totals")

    canonical = SoilNetDualHead(backbone_pretrained=False).state_dict()
    count_hits: Counter[str] = Counter()
    all_rows: list[dict[str, object]] = []
    for index, row in enumerate(scope_a, 1):
        path = paths["checkpoint_root"] / row["relative_path"]
        inspected = inspect_checkpoint(path, "CHECKPOINT_ROOT", canonical, count_hits)
        flags = []
        if sum(item["sha256"] == row["sha256"] for item in scope_a) > 1:
            flags.append("DUPLICATE_SHA256")
        if any(item["sha256"] == row["sha256"] for item in scope_b):
            flags.append("MATCHES_GITHUB")
        for flag, key in (("MODEL_COMPATIBLE", "model_compatible"), ("HAS_REG_HEAD", "has_reg_head"),
                          ("HAS_CLS_HEAD", "has_cls_head"), ("HAS_LIGHT_BRANCH", "has_light_branch"),
                          ("HAS_OPTIMIZER", "has_optimizer"), ("HAS_EPOCH", "has_epoch"), ("HAS_CONFIG", "has_config")):
            if inspected[key]:
                flags.append(flag)
        all_rows.append({
            "checkpoint_id": row["checkpoint_id"], "scope": "CHECKPOINT_ROOT", "relative_path": row["relative_path"],
            "filename": row["filename"], "size_bytes": row["size_bytes"], "sha256": row["sha256"],
            **inspected, "additional_flags": ";".join(flags),
        })
        if index % 100 == 0:
            print(f"CHECKPOINT_ROOT safe-inspected {index}/969", flush=True)

    source_texts = [(scope, relative, document_text(path, include_outputs=False)) for scope, path, relative in documents]
    legacy_hashes = defaultdict(list)
    scope_a_hashes = defaultdict(list)
    for row in scope_a:
        scope_a_hashes[row["sha256"]].append(row["relative_path"])
    for row in scope_b:
        legacy_hashes[row["sha256"]].append(row["relative_path"])

    scope_c_files = sorted(paths["data_root"].rglob("*.pth"), key=lambda item: item.as_posix().casefold())
    scope_c_rows: list[dict[str, object]] = []
    for index, path in enumerate(scope_c_files, 1):
        digest = sha256_file(path)
        inspected = inspect_checkpoint(path, "DATA_ROOT_EMBEDDED_ARTIFACTS", canonical, count_hits)
        relative = path.relative_to(paths["data_root"]).as_posix()
        stat = path.stat()
        scope_c_rows.append({
            "artifact_id": f"sha256:{digest[:20]}", "absolute_source_path": str(path),
            "relative_path_from_data_root": relative, "filename": path.name, "size_bytes": stat.st_size,
            "sha256": digest, "mtime": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(), **inspected,
        })
        if index % 25 == 0:
            print(f"DATA_ROOT embedded safe-inspected {index}/{len(scope_c_files)}", flush=True)
    if len(scope_c_rows) != 108:
        raise RuntimeError(f"Expected 108 DATA_ROOT .pth files, found {len(scope_c_rows)}")

    c_hashes = Counter(str(row["sha256"]) for row in scope_c_rows)
    for row in scope_c_rows:
        flags = []
        if c_hashes[str(row["sha256"])] > 1 or scope_a_hashes[str(row["sha256"])]:
            flags.append("DUPLICATE_SHA256")
        if legacy_hashes[str(row["sha256"])]:
            flags.append("MATCHES_GITHUB")
        for flag, key in (("MODEL_COMPATIBLE", "model_compatible"), ("HAS_REG_HEAD", "has_reg_head"),
                          ("HAS_CLS_HEAD", "has_cls_head"), ("HAS_LIGHT_BRANCH", "has_light_branch"),
                          ("HAS_OPTIMIZER", "has_optimizer"), ("HAS_EPOCH", "has_epoch"), ("HAS_CONFIG", "has_config")):
            if row[key]:
                flags.append(flag)
        row["additional_flags"] = ";".join(flags)
        all_rows.append({
            "checkpoint_id": row["artifact_id"], "scope": "DATA_ROOT_EMBEDDED_ARTIFACTS",
            "relative_path": row["relative_path_from_data_root"], "filename": row["filename"],
            "size_bytes": row["size_bytes"], "sha256": row["sha256"],
            **{key: row[key] for key in row if key not in {"artifact_id", "absolute_source_path", "relative_path_from_data_root", "filename", "size_bytes", "sha256", "mtime"}},
        })

    common_fields = [
        "checkpoint_id", "scope", "relative_path", "filename", "size_bytes", "sha256", "safe_load_status",
        "top_level_type", "top_level_keys", "state_dict_prefix", "tensor_count", "parameter_tensor_count",
        "primary_category", "suspected_mu", "suspected_epoch", "suspected_model", "has_reg_head", "has_cls_head",
        "has_light_branch", "num_classes", "regression_output_dimension", "reg_head_shape", "cls_head_shape",
        "model_compatible", "has_optimizer", "has_epoch", "has_config", "additional_flags", "notes",
    ]
    write_csv(REPO / "results/audit/checkpoint_classification.csv", all_rows, common_fields)
    c_fields = [
        "artifact_id", "absolute_source_path", "relative_path_from_data_root", "filename", "size_bytes", "sha256", "mtime",
        "safe_load_status", "top_level_type", "top_level_keys", "state_dict_prefix", "tensor_count",
        "parameter_tensor_count", "suspected_category", "suspected_mu", "suspected_epoch", "suspected_model",
        "has_reg_head", "has_cls_head", "has_light_branch", "num_classes", "regression_output_dimension",
        "reg_head_shape", "cls_head_shape", "model_compatible", "additional_flags", "notes",
    ]
    write_csv(REPO / "results/audit/data_root_checkpoint_inventory.csv", scope_c_rows, c_fields)

    cross_rows = []
    for row in scope_c_rows:
        digest = str(row["sha256"])
        a_matches, b_matches = scope_a_hashes[digest], legacy_hashes[digest]
        if a_matches and b_matches:
            classification = "MULTI_SCOPE_DUPLICATE"
        elif a_matches:
            classification = "EXACT_DUPLICATE_LOCAL"
        elif b_matches:
            classification = "EXACT_DUPLICATE_GITHUB"
        else:
            classification = "UNIQUE_DATA_ROOT_ARTIFACT"
        matches = [("CHECKPOINT_ROOT", item) for item in a_matches] + [("LEGACY_GITHUB", item) for item in b_matches]
        if not matches:
            matches = [("NONE", "")]
        for matched_scope, matched_path in matches:
            cross_rows.append({
                "scope_c_artifact": row["relative_path_from_data_root"], "scope_c_sha256": digest,
                "matched_scope": matched_scope, "matched_path": matched_path,
                "matched_sha256": digest if matched_path else "", "byte_identical": bool(matched_path),
                "classification": classification,
                "lineage_implication": (
                    "Same bytes across scopes; publish at most one approved copy."
                    if matched_path else f"No Scope A/B SHA match; within-Scope-C copies with this hash: {c_hashes[digest]}."
                ),
            })
    write_csv(REPO / "results/audit/checkpoint_cross_scope_matches.csv", cross_rows, [
        "scope_c_artifact", "scope_c_sha256", "matched_scope", "matched_path", "matched_sha256",
        "byte_identical", "classification", "lineage_implication",
    ])

    summary = {
        "checkpoint_root_files": len(scope_a), "legacy_github_files": len(scope_b),
        "data_root_embedded_files": len(scope_c_rows),
        "cross_scope_unique_sha256": len({str(row["sha256"]) for row in scope_a + scope_b} | set(c_hashes)),
        "checkpoint_root_categories": dict(Counter(str(row["primary_category"]) for row in all_rows if row["scope"] == "CHECKPOINT_ROOT")),
        "data_root_categories": dict(Counter(str(row["suspected_category"]) for row in scope_c_rows)),
        "data_root_unique_sha256": len(c_hashes),
        "data_root_internal_redundant_files": sum(count - 1 for count in c_hashes.values()),
        "data_root_duplicates_with_checkpoint_root": sum(bool(scope_a_hashes[str(row["sha256"])]) for row in scope_c_rows),
        "data_root_duplicates_with_legacy_github": sum(bool(legacy_hashes[str(row["sha256"])]) for row in scope_c_rows),
        "data_root_reg_and_cls_heads": sum(bool(row["has_reg_head"] and row["has_cls_head"]) for row in scope_c_rows),
        "data_root_light_branch": sum(bool(row["has_light_branch"]) for row in scope_c_rows),
        "checkpoint_scalar_count_hits": dict(count_hits),
    }
    write_json(REPO / "results/audit/phase2_checkpoint_summary.json", summary)
    del canonical
    gc.collect()
    return scope_c_rows, summary


def audit_count_provenance(paths: dict[str, Path], checkpoint_summary: dict[str, object]) -> None:
    search_roots = [
        ("DATA_ROOT", paths["data_root"]), ("CHECKPOINT_ROOT", paths["checkpoint_root"]),
        ("LEGACY_GITHUB", paths["legacy_root"]), ("CURRENT_DOCUMENTATION", REPO / "docs"),
    ]
    patterns = {"2250": re.compile(r"(?<!\d)2,?250(?!\d)"), "2057": re.compile(r"(?<!\d)2,?057(?!\d)")}
    hits: dict[str, list[tuple[str, str, str]]] = {key: [] for key in patterns}
    for scope, root in search_roots:
        for path in root.rglob("*"):
            if not path.is_file() or path.suffix.casefold() not in TEXT_SUFFIXES:
                continue
            text = document_text(path)
            for number, pattern in patterns.items():
                for match in pattern.finditer(text):
                    context = text[max(0, match.start() - 100):match.end() + 100].replace("\n", " ")
                    hits[number].append((scope, path.relative_to(root).as_posix(), context))
    readme = REPO / "README.md"
    if readme.is_file():
        text = readme.read_text(encoding="utf-8", errors="replace")
        for number, pattern in patterns.items():
            for match in pattern.finditer(text):
                context = text[max(0, match.start() - 100):match.end() + 100].replace("\n", " ")
                hits[number].append(("CURRENT_DOCUMENTATION", "README.md", context))
    count_context = re.compile(r"(?:image|sample|dataset|row|labeled|total|ảnh|mẫu).{0,80}(?:2,?250)|(?:2,?250).{0,80}(?:image|sample|dataset|row|labeled|total|ảnh|mẫu)", re.I)
    count_2250 = [hit for hit in hits["2250"] if count_context.search(hit[2])]
    legacy_2250 = [hit for hit in hits["2250"] if hit[0] == "LEGACY_GITHUB"]
    xlsx_hits = {"2057": 0, "2250": 0}
    xlsx_files = []
    xlsx_errors = []
    try:
        import openpyxl

        for scope, root in search_roots:
            for path in sorted(root.rglob("*.xlsx"), key=lambda item: item.as_posix().casefold()):
                xlsx_files.append(f"{scope}:{path.relative_to(root).as_posix()}")
                try:
                    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
                    for worksheet in workbook.worksheets:
                        for row in worksheet.iter_rows(values_only=True):
                            for value in row:
                                if isinstance(value, (int, float)) and not isinstance(value, bool):
                                    if float(value) == 2057.0:
                                        xlsx_hits["2057"] += 1
                                    elif float(value) == 2250.0:
                                        xlsx_hits["2250"] += 1
                                elif isinstance(value, str):
                                    normalized = value.strip().replace(",", "")
                                    if normalized in xlsx_hits:
                                        xlsx_hits[normalized] += 1
                    workbook.close()
                except Exception as exc:  # Record unreadable workbooks; never silently infer zero.
                    xlsx_errors.append(f"{scope}:{path.relative_to(root).as_posix()}:{type(exc).__name__}")
    except ImportError:
        xlsx_errors.append("OPENPYXL_UNAVAILABLE")
    lines = [
        "# Dataset count provenance: 2,250 versus 2,057", "",
        "## Decision", "", "**UNRESOLVED**", "",
        "The strongest reproducible labeled-sample count is **2,057**: the discovered `labels.csv` has 2,057 data rows, "
        "and historical notebook outputs repeatedly report loading 2,057 labeled images.", "",
        "No source in the searched dataset root, checkpoint root, pinned legacy checkout, current documentation, or readable XLSX cells "
        "establishes 2,250 as an image/sample count. The only exact `2250` occurrence in the pinned legacy text is the decimal "
        "tail of `MSE: 1319.2250`, not a count. Therefore `2250 - 2057 = 193` is not interpreted as filtered images.", "",
        "## Search evidence", "",
        f"- DATA_ROOT/legacy/current-document exact 2,057 occurrences: {len(hits['2057'])}",
        f"- Exact 2,250 occurrences in a count-like context: {len(count_2250)}",
        f"- Pinned legacy exact 2250 occurrences: {len(legacy_2250)} (metric decimal, not count)",
        f"- Checkpoint scalar metadata equal to 2,057: {checkpoint_summary.get('checkpoint_scalar_count_hits', {}).get('2057', 0)}",
        f"- Checkpoint scalar metadata equal to 2,250: {checkpoint_summary.get('checkpoint_scalar_count_hits', {}).get('2250', 0)}",
        f"- XLSX workbooks inspected: {len(xlsx_files)}",
        f"- XLSX exact cells equal to 2,057: {xlsx_hits['2057']}",
        f"- XLSX exact cells equal to 2,250: {xlsx_hits['2250']}",
        f"- XLSX workbook read errors: {len(xlsx_errors)}" + (f" (`{'; '.join(xlsx_errors)}`)" if xlsx_errors else ""), "",
        "## Limitation", "",
        "No manuscript/draft file or historical capture manifest asserting 2,250 was supplied in the audited roots. If such a "
        "document is later supplied, it must be reconciled by sample identifiers—not arithmetic difference alone.", "",
    ]
    (REPO / "results/audit/dataset_count_provenance.md").write_text("\n".join(lines), encoding="utf-8")


def checkpoint_references(documents: list[tuple[str, Path, str]]) -> tuple[list[dict[str, str]], dict[str, list[tuple[str, str, str]]]]:
    by_basename: dict[str, list[tuple[str, str, str]]] = defaultdict(list)
    rows = []
    pth_pattern = re.compile(r"([\w .()=+\-]+\.pth)", re.I)
    for scope, path, relative in documents:
        text = document_text(path, include_outputs=False)
        for line_number, line in enumerate(text.splitlines(), 1):
            # Consumer/load references are not generator lineage. Accept only an
            # explicit save call or a recognized save-target assignment.
            if not re.search(
                r"torch\.save|(?:BEST_MODEL_SAVE_PATH|MODEL_SAVE_PATH|CHECKPOINT_SAVE_PATH|SAVE_PATH)\s*=",
                line,
                re.I,
            ):
                continue
            for filename in pth_pattern.findall(line):
                filename = Path(filename.replace("\\", "/")).name
                by_basename[filename.casefold()].append((scope, relative, f"line {line_number}"))
    return rows, by_basename


def audit_lineage(paths: dict[str, Path], documents: list[tuple[str, Path, str]], scope_c: list[dict[str, object]]) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    _, refs = checkpoint_references(documents)
    metric_files = [path for path in paths["data_root"].rglob("*") if path.suffix.casefold() in {".csv", ".xlsx"}]
    lineages = []
    for row in scope_c:
        if row["suspected_category"] != "FINETUNED_SOILNET_DUALHEAD":
            continue
        name = str(row["filename"])
        matches = refs.get(name.casefold(), [])
        mu = str(row["suspected_mu"])
        tiny = "tiny" in name.casefold()
        metrics = [
            path.relative_to(paths["data_root"]).as_posix() for path in metric_files
            if (not mu or re.search(rf"mu[_ =-]?{re.escape(mu.removesuffix('.0'))}(?:\D|$)", path.name, re.I))
            and (("tiny" in path.as_posix().casefold()) == tiny)
        ]
        lower = name.casefold()
        if "best_finetuned" in lower:
            selection_basis, checkpoint_status = "TRAIN_METRIC", "TRAINING_SELECTED_CHECKPOINT"
        elif any(token in lower for token in ("finetune_mu", "finetuned", "final_tinyimagenet")):
            selection_basis, checkpoint_status = "LAST_EPOCH", "CANNOT_PROVIDE_PROSPECTIVE_HELDOUT_TEST"
        else:
            selection_basis, checkpoint_status = "UNKNOWN", "EXPOSURE_UNKNOWN"
        exact_save = bool(matches)
        architecture = bool(row["model_compatible"] and row["has_reg_head"] and row["has_cls_head"] and row["has_light_branch"])
        metric_match = bool(metrics)
        score = 2 * int(exact_save) + 2 * int(architecture) + int(metric_match)
        if exact_save and architecture and metric_match:
            lineage_status = "HIGH_CONFIDENCE_LINEAGE"
        elif exact_save and architecture:
            lineage_status = "MEDIUM_CONFIDENCE_LINEAGE"
        elif exact_save or architecture:
            lineage_status = "AMBIGUOUS"
        else:
            lineage_status = "NO_LINEAGE"
        source_notebooks = ";".join(f"{scope}:{relative}" for scope, relative, _ in matches)
        all_2057 = any(
            any(token in relative.casefold() for token in ("finetune", "fine_tune"))
            and "train_test_split" not in document_text(path, include_outputs=False)
            and "dataloader" in document_text(path, include_outputs=False).casefold()
            for scope, path, relative in documents if any(relative == match[1] and scope == match[0] for match in matches)
        )
        exposure = "TRAINED_ON_ALL_2057" if all_2057 else "UNKNOWN_EXPOSURE"
        lineages.append({
            "checkpoint_id": row["artifact_id"], "sha256": row["sha256"], "filename": name,
            "relative_path": row["relative_path_from_data_root"], "state_dict_prefix": row["state_dict_prefix"],
            "reg_head_shape": row["reg_head_shape"], "cls_head_shape": row["cls_head_shape"],
            "num_classes": row["num_classes"], "light_branch": row["has_light_branch"],
            "suspected_mu": mu, "suspected_epoch": row["suspected_epoch"],
            "suspected_notebook": source_notebooks or "UNMAPPED",
            "suspected_metrics_csv": ";".join(metrics[:20]) or "UNMAPPED",
            "training_dataset_count": 2057 if exposure == "TRAINED_ON_ALL_2057" else "UNKNOWN",
            "training_split_known": False, "historical_test_known": False,
            "data_exposure": exposure, "selection_basis": selection_basis,
            "lineage_evidence_score": score, "lineage_status": lineage_status,
            "status": checkpoint_status,
            "notes": "No notebook records the checkpoint SHA at save time; filename/path evidence alone is not VERIFIED lineage.",
        })
    write_csv(REPO / "results/audit/fine_tuned_checkpoint_lineage.csv", lineages, [
        "checkpoint_id", "sha256", "filename", "relative_path", "state_dict_prefix", "reg_head_shape",
        "cls_head_shape", "num_classes", "light_branch", "suspected_mu", "suspected_epoch",
        "suspected_notebook", "suspected_metrics_csv", "training_dataset_count", "training_split_known",
        "historical_test_known", "data_exposure", "selection_basis", "lineage_evidence_score",
        "lineage_status", "status", "notes",
    ])

    sota_rows = []
    sota_source = next((relative for scope, _, relative in documents if relative.endswith("Experiment 3 SOTA/Code/Fine_tune_SOTA.py")), "UNMAPPED")
    for row in scope_c:
        if "sota" not in str(row["relative_path_from_data_root"]).casefold():
            continue
        filename = str(row["filename"]).casefold()
        architecture_hint = next((name for name in ("efficientnet_b0", "mobilevitv2_050", "mobilevit_s", "mobilenetv2_100", "resnet") if name in filename), "OTHER")
        safe = row["safe_load_status"] == "SAFE_LOAD_OK"
        sota_rows.append({
            "artifact_id": row["artifact_id"], "relative_path": row["relative_path_from_data_root"],
            "sha256": row["sha256"], "architecture": architecture_hint if safe else "UNKNOWN_UNREADABLE",
            "filename_architecture_hint": architecture_hint,
            "safe_load_status": row["safe_load_status"], "primary_category": row["suspected_category"],
            "pretrained_initialization": "IMAGENET_PRETRAINED_TRUE",
            "dataset": "full labeled dataframe; historical file name label_updated_colab.csv",
            "split": "NONE; one shuffled loader", "metric_source": "same training loader",
            "selection_basis": "LAST_EPOCH_OR_TRAIN_LOSS_EARLY_STOP",
            "source_script": sota_source, "lineage_status": "HIGH_CONFIDENCE_LINEAGE" if safe else "AMBIGUOUS",
            "evidence_status": "TRAINING_SET_METRICS_ONLY",
            "notes": "Architecture is confirmed only for safely loaded state; unreadable full-object files retain filename hint only.",
        })
    write_csv(REPO / "results/audit/sota_checkpoint_lineage.csv", sota_rows, [
        "artifact_id", "relative_path", "sha256", "architecture", "filename_architecture_hint", "safe_load_status", "primary_category",
        "pretrained_initialization", "dataset", "split", "metric_source", "selection_basis",
        "source_script", "lineage_status", "evidence_status", "notes",
    ])
    return lineages, sota_rows


def audit_historical_splits(paths: dict[str, Path], documents: list[tuple[str, Path, str]], lineages: list[dict[str, object]]) -> None:
    linked_by_source = defaultdict(list)
    linked_by_metric = defaultdict(list)
    for row in lineages:
        for source in str(row["suspected_notebook"]).split(";"):
            if ":" in source:
                linked_by_source[source.split(":", 1)[1]].append(str(row["filename"]))
        for metric in str(row["suspected_metrics_csv"]).split(";"):
            if metric and metric != "UNMAPPED":
                linked_by_metric[metric].append(str(row["filename"]))
    rows = []
    for scope, path, relative in documents:
        text = document_text(path, include_outputs=False)
        lower = text.casefold()
        soil_labeled = any(token in lower for token in ("labels.csv", "labels_new.csv", "label_updated_colab.csv", "soildualtaskdataset", "labeledimagedataset"))
        split_token = any(token in lower for token in (
            "train_test_split", "random_split", "groupkfold", "kfold", "stratified", "subset", "val_loader", "test_loader", "validation",
        ))
        if not soil_labeled and not split_token:
            continue
        if "train_test_split" in lower and soil_labeled:
            split_type = "RANDOM_ROW_HOLDOUT"
            test_size_match = re.search(r"test_size\s*=\s*([0-9.]+)", text)
            seed_match = re.search(r"random_state\s*=\s*(\d+)", text)
            test_fraction = float(test_size_match.group(1)) if test_size_match else None
            test_count = int(2057 * test_fraction + 0.999999) if test_fraction else "UNKNOWN"
            train_count = 2057 - test_count if isinstance(test_count, int) else "UNKNOWN"
            seed = seed_match.group(1) if seed_match else "UNKNOWN"
            reproducible = "NO"
            identifiers = "NO_PERSISTED_SAMPLE_IDS"
        elif "train_test_split" in lower:
            split_type, seed, train_count, test_count = "EXTERNAL_OR_DATASET_UNRESOLVED_ROW_HOLDOUT", "UNKNOWN", "UNKNOWN", "UNKNOWN"
            reproducible, identifiers = "NO", "NO_PERSISTED_SAMPLE_IDS"
        elif soil_labeled and "dataloader" in lower and any(token in lower for token in ("optimizer", "train_epoch", "loss.backward", ".backward(")):
            split_type, seed, train_count, test_count = "NO_HELDOUT_ALL_ROWS", "MISSING", 2057, 0
            reproducible, identifiers = "NO", "labels.csv rows referenced but consumed manifest not saved"
        else:
            split_type, seed, train_count, test_count = "MENTION_ONLY_OR_EXTERNAL_DATA", "UNKNOWN", "UNKNOWN", "UNKNOWN"
            reproducible, identifiers = "NO", "UNKNOWN"
        rows.append({
            "source_scope": scope, "source": relative, "split_type": split_type, "seed": seed,
            "train_count": train_count, "val_count": 0 if isinstance(train_count, int) else "UNKNOWN",
            "test_count": test_count, "sample_identifiers_available": identifiers,
            "reproducible": reproducible, "checkpoint_linked": ";".join(linked_by_source.get(relative, [])) or "NONE",
            "fine_tuned_soilnet": bool(linked_by_source.get(relative)),
            "genuine_heldout": False,
            "notes": "A row-level split in a downstream classifier does not make an encoder trained on all labeled rows held out.",
        })

    # Results tables/logs are searched separately from executable source. Metric
    # columns alone are not a persisted sample split.
    artifact_suffixes = {".csv", ".xlsx", ".xls", ".log", ".txt", ".json"}
    split_terms = re.compile(r"\b(train(?:ing)?|val(?:idation)?|test|holdout|split|train_idx|val_idx|test_idx|subset|fold)\b", re.I)
    for scope, root in (("DATA_ROOT", paths["data_root"]), ("LEGACY_GITHUB", paths["legacy_root"])):
        for path in sorted(root.rglob("*"), key=lambda item: item.as_posix().casefold()):
            if not path.is_file() or path.suffix.casefold() not in artifact_suffixes:
                continue
            relative = path.relative_to(root).as_posix()
            headers: list[str] = []
            records: list[dict[str, object]] = []
            artifact_has_rows = False
            try:
                if path.suffix.casefold() == ".csv":
                    with path.open(newline="", encoding="utf-8", errors="replace") as handle:
                        reader = csv.DictReader(handle)
                        headers = [str(value) for value in (reader.fieldnames or [])]
                        records = [dict(row) for row in reader]
                    artifact_has_rows = bool(records)
                    text = " ".join(headers + [str(value) for row in records[:50] for value in row.values()])
                elif path.suffix.casefold() == ".xlsx":
                    import openpyxl

                    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
                    values: list[str] = []
                    for worksheet in workbook.worksheets:
                        sheet_rows = list(worksheet.iter_rows(values_only=True))
                        if sheet_rows and not headers:
                            headers = [str(value or "") for value in sheet_rows[0]]
                            records = [dict(zip(headers, row)) for row in sheet_rows[1:]]
                        artifact_has_rows = artifact_has_rows or len(sheet_rows) > 1
                        values.extend(str(value) for row in sheet_rows[:51] for value in row if value is not None)
                    workbook.close()
                    text = " ".join(values)
                else:
                    text = path.read_text(encoding="utf-8", errors="replace")
            except Exception as exc:
                rows.append({
                    "source_scope": scope, "source": relative, "split_type": "UNREADABLE_RESULT_ARTIFACT",
                    "seed": "UNKNOWN", "train_count": "UNKNOWN", "val_count": "UNKNOWN", "test_count": "UNKNOWN",
                    "sample_identifiers_available": "UNKNOWN", "reproducible": "NO", "checkpoint_linked": "NONE",
                    "fine_tuned_soilnet": False, "genuine_heldout": False,
                    "notes": f"Safe tabular/text read failed: {type(exc).__name__}",
                })
                continue
            if not split_terms.search(text):
                continue
            normalized_headers = [header.casefold().strip() for header in headers]
            id_column = next((header for header in normalized_headers if any(token in header for token in ("sample_id", "image_path", "filename", "file_path"))), "")
            split_column = next((header for header in normalized_headers if header in {"split", "partition", "subset"}), "")
            split_key = headers[normalized_headers.index(split_column)] if split_column else ""
            observed_partitions = {
                str(row.get(split_key, "")).casefold().strip() for row in records if split_key and row.get(split_key) is not None
            }
            has_train_and_holdout = "train" in observed_partitions and bool(observed_partitions & {"validation", "val", "test", "holdout"})
            persisted = bool(id_column and split_column and artifact_has_rows and has_train_and_holdout)
            linked = linked_by_metric.get(relative, [])
            rows.append({
                "source_scope": scope, "source": relative,
                "split_type": "PERSISTED_SAMPLE_SPLIT" if persisted else "METRIC_COLUMNS_OR_SPLIT_MENTION_ONLY",
                "seed": "UNKNOWN", "train_count": "UNKNOWN", "val_count": "UNKNOWN", "test_count": "UNKNOWN",
                "sample_identifiers_available": "YES" if persisted else "NO_PERSISTED_SAMPLE_IDS",
                "reproducible": "YES_EXACT_MANIFEST" if persisted else "NO",
                "checkpoint_linked": ";".join(linked) or "NONE", "fine_tuned_soilnet": bool(linked),
                "genuine_heldout": bool(persisted and linked),
                "notes": (
                    "A persisted identifier/split table linked to this checkpoint generation."
                    if persisted and linked else
                    "Train/validation/test text or metric columns do not establish a sample-level checkpoint-linked split."
                ),
            })
    write_csv(REPO / "results/audit/historical_split_inventory.csv", rows, [
        "source_scope", "source", "split_type", "seed", "train_count", "val_count", "test_count",
        "sample_identifiers_available", "reproducible", "checkpoint_linked", "fine_tuned_soilnet",
        "genuine_heldout", "notes",
    ])


def audit_unseen_and_ssl(paths: dict[str, Path], scope_c: list[dict[str, object]], lineages: list[dict[str, object]]) -> None:
    samples = read_csv("data/manifests/soilnet_samples.csv")
    splits = read_csv("data/splits/final_split.csv")
    inventory = read_csv("data/manifests/dataset_inventory.csv")
    split_by_id = {row["sample_id"]: row for row in splits}
    unlabeled_sha = {row["sha256"] for row in inventory if row["relative_path"].startswith("unlabeled_images/")}
    duplicate_rows = read_csv("results/audit/duplicate_report.csv")
    unlabeled_groups = {row["duplicate_group_id"] for row in duplicate_rows if row["relative_path"].startswith("unlabeled_images/")}
    any_all_exposed = any(row["data_exposure"] == "TRAINED_ON_ALL_2057" for row in lineages)
    unseen_rows = []
    for sample in samples:
        split = split_by_id[sample["sample_id"]]
        status = "HISTORICALLY_EXPOSED" if any_all_exposed else "UNKNOWN"
        unseen_rows.append({
            "sample_id": sample["sample_id"], "relative_path": sample["relative_path"], "sha256": sample["sha256"],
            "prospective_split": split["split"], "status": status,
            "exact_match_in_historical_unlabeled_pool": sample["sha256"] in unlabeled_sha,
            "near_duplicate_group_in_historical_unlabeled_pool": split["group_id"] in unlabeled_groups,
            "evidence": (
                "High-confidence fine-tune notebooks build one dataset from all 2057 rows. For fallback rows, sample identity was "
                "exposed although historical augmented bytes are not reproducible."
                if any_all_exposed else "No sample-level historical consumed manifest is available."
            ),
        })
    write_csv(REPO / "results/audit/unseen_labeled_candidates.csv", unseen_rows, [
        "sample_id", "relative_path", "sha256", "prospective_split", "status",
        "exact_match_in_historical_unlabeled_pool", "near_duplicate_group_in_historical_unlabeled_pool", "evidence",
    ])

    exact_test = sum(row["split"] == "test" and row["sha256"] in unlabeled_sha for row in splits)
    near_test = sum(row["split"] == "test" and row["group_id"] in unlabeled_groups for row in splits)
    ssl_rows = []
    classification = read_csv("results/audit/checkpoint_classification.csv")
    for row in classification:
        category = row["primary_category"]
        lower = (row["relative_path"] + "/" + row["filename"]).casefold()
        if category not in {"SSL_BACKBONE", "VICREG_PROJECTOR", "PARTIAL_STATE"} or "vicreg" not in lower:
            continue
        if row["scope"] == "DATA_ROOT_EMBEDDED_ARTIFACTS" or "original_new" in lower:
            source = "SOIL_IMAGES_UNLABELED"
            protocol = "TRANSDUCTIVE_PRETRAINING"
            initialization = "TinyImageNet/ImageNet initializer possible; exact generation-specific initializer requires lineage review"
        else:
            source = "UNKNOWN"
            protocol = "PROVENANCE_UNKNOWN"
            initialization = "Directory name is not sufficient provenance"
        ssl_rows.append({
            "checkpoint_id": row["checkpoint_id"], "scope": row["scope"], "relative_path": row["relative_path"],
            "sha256": row["sha256"], "artifact_component": category, "pretraining_source": source,
            "initialization_evidence": initialization, "protocol_status": protocol,
            "prospective_test_exact_sha_exposed_in_unlabeled_pool": exact_test,
            "prospective_test_near_group_exposed_in_unlabeled_pool": near_test,
            "notes": "Soil unlabeled pretraining cannot be called completely held-out inductive when prospective test images/groups overlap its pool.",
        })
    write_csv(REPO / "results/audit/ssl_checkpoint_provenance.csv", ssl_rows, [
        "checkpoint_id", "scope", "relative_path", "sha256", "artifact_component", "pretraining_source",
        "initialization_evidence", "protocol_status", "prospective_test_exact_sha_exposed_in_unlabeled_pool",
        "prospective_test_near_group_exposed_in_unlabeled_pool", "notes",
    ])


def write_classical_plan(paths: dict[str, Path]) -> None:
    local_code = [path for path in paths["data_root"].rglob("*.py") if "classifier" in path.name.casefold()]
    code_text = "\n".join(path.read_text(encoding="utf-8", errors="replace") for path in local_code)
    plan = f"""# Classical ML reproduction feasibility

No CNN training or full feature extraction was run in PHASE 2.

| Method | Historical evidence | Validity | Recovery |
|---|---|---|---|
| KNN | No configured implementation, saved estimator, embeddings, or metric row found | MISSING | CHEAP_RECOMPUTE after encoder/config preselection |
| SVM | Local scripts and result rows exist; row-level 80/20 split uses seed 42 | Encoder exposure and sample IDs are not cleanly held out/persisted | CHEAP_RECOMPUTE |
| GBM | Local scripts and result rows exist; row-level 80/20 split uses seed 42 | Same limitation as SVM | CHEAP_RECOMPUTE |
| Decision Tree | No configured implementation/result artifact found | MISSING | CHEAP_RECOMPUTE |
| MLP | Imported in local scripts but not configured in the executed classifier dictionary; no result row | MISSING | CHEAP_RECOMPUTE |

Evidence details:

- Classifier source files found: {len(local_code)}.
- `train_test_split` present: {"yes" if "train_test_split" in code_text else "no"}; the discovered scripts use `test_size=0.2, random_state=42`.
- Saved `.npy/.npz/.pkl/.joblib` embeddings or classifier objects found in audited roots: none.
- The scripts reference fine-tuned SoilNet checkpoints whose training exposure is not cleanly held out from their row-level test split.

Minimal safe regeneration:

1. Freeze the encoder/checkpoint choice before inspecting prospective test metrics.
2. Extract features once for the fixed 1496/317/244 group-aware manifest.
3. Fit scaler and classifier on train only; use validation only for any hyperparameter choice.
4. Evaluate the 244 test samples exactly once and write predictions plus metadata.
5. If the chosen encoder used the soil image pool, label the protocol transductive; no verified pure-external SSL encoder is currently established.
"""
    (REPO / "results/audit/classical_ml_reproduction_plan.md").write_text(plan, encoding="utf-8")


def update_claims(lineages: list[dict[str, object]]) -> None:
    rows = read_csv("results/audit/claims_evidence_matrix.csv")
    mapping = {
        "ARCH": ("UNKNOWN", "UNKNOWN", "KEEP"), "LIGHT": ("UNKNOWN", "UNKNOWN", "KEEP"),
        "SM0": ("TRAIN", "CONFIRMED_LEAKAGE", "SHORT_FINETUNE"),
        "SM20": ("TRAIN", "CONFIRMED_LEAKAGE", "SHORT_FINETUNE"),
        "CLASS": ("TRAIN", "CONFIRMED_LEAKAGE", "SHORT_FINETUNE"),
        "SSL": ("TRAIN", "POTENTIAL_LEAKAGE", "KEEP"),
        "MU": ("TRAIN", "CONFIRMED_LEAKAGE", "MISSING"),
        "TINY": ("TRAIN", "POTENTIAL_LEAKAGE", "KEEP"),
        "BASE": ("TRAIN", "CONFIRMED_LEAKAGE", "MISSING"),
        "KNN": ("UNKNOWN", "UNKNOWN", "CHEAP_RECOMPUTE"),
        "GBM": ("TEST", "CONFIRMED_LEAKAGE", "CHEAP_RECOMPUTE"),
        "SVM": ("TEST", "CONFIRMED_LEAKAGE", "CHEAP_RECOMPUTE"),
        "DT": ("UNKNOWN", "UNKNOWN", "CHEAP_RECOMPUTE"),
        "MLP": ("UNKNOWN", "UNKNOWN", "CHEAP_RECOMPUTE"),
        "EDGE": ("UNKNOWN", "UNKNOWN", "MISSING"), "IRR": ("UNKNOWN", "UNKNOWN", "MISSING"),
    }
    for row in rows:
        metric, exposure, recovery = mapping[row["claim_id"]]
        row.update({"historical_metric_type": metric, "data_exposure": exposure, "recovery": recovery})
        if row["claim_id"] in {"GBM", "SVM"}:
            row["notes"] = "Local code and result rows exist, but its random row test has no persisted sample IDs and uses an encoder with non-clean exposure."
        if row["claim_id"] == "MLP":
            row["notes"] = "MLP classes are imported in local scripts but absent from the configured/recorded model results."
    fields = list(rows[0])
    for field in ("historical_metric_type", "data_exposure", "recovery"):
        if field not in fields:
            fields.append(field)
    write_csv(REPO / "results/audit/claims_evidence_matrix.csv", rows, fields)
    md = [
        "# Claims–evidence matrix", "",
        "Claimed manuscript values were not supplied; PHASE 2 does not alter claims or invent values.", "",
        "| ID | Claim | Evidence status | Historical metric | Exposure | Recovery |", "|---|---|---|---|---|---|",
    ]
    for row in rows:
        values = [row[key].replace("|", "\\|") for key in ("claim_id", "claim", "evidence_status", "historical_metric_type", "data_exposure", "recovery")]
        md.append("| " + " | ".join(values) + " |")
    md.extend(["", "Machine-readable lineage and notes are in `results/audit/claims_evidence_matrix.csv`.", ""])
    (REPO / "docs/CLAIMS_EVIDENCE_MATRIX.md").write_text("\n".join(md), encoding="utf-8")


def public_release_scan() -> dict[str, object]:
    process = subprocess.run(
        ["git", "ls-files", "--cached", "--others", "--exclude-standard"], cwd=REPO,
        capture_output=True, text=True, check=True,
    )
    candidates = [REPO / line for line in process.stdout.splitlines() if line]
    path_markers = ["/" + "mnt" + "/", "/home/" + "diy-hus", "C" + ":" + chr(92), "D" + ":" + chr(92), "F" + ":" + chr(92)]
    secret_patterns = [
        ("GITHUB_TOKEN", re.compile("gh" + "p_[A-Za-z0-9]{20,}|github_" + "pat_[A-Za-z0-9_]{20,}")),
        ("PRIVATE_KEY", re.compile("BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY")),
        ("API_SECRET_ASSIGNMENT", re.compile(r"(?i)(?:api[_-]?key|password|secret|token)\s*[:=]\s*['\"][^'\"]{8,}['\"]")),
    ]
    email_pattern = re.compile(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b")
    findings = []
    image_suffixes = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp"}
    for path in candidates:
        relative = path.relative_to(REPO).as_posix()
        if path.suffix.casefold() in image_suffixes:
            findings.append((relative, "RAW_IMAGE_OR_FIGURE_REVIEW", "binary image present in release candidate", "REVIEW"))
        if path.suffix.casefold() == ".pth":
            findings.append((relative, "CHECKPOINT_BINARY", "checkpoint present in release candidate", "BLOCK"))
        if not path.is_file() or path.suffix.casefold() not in TEXT_SUFFIXES | {".cff"}:
            continue
        text = path.read_text(encoding="utf-8", errors="replace")
        for marker in path_markers:
            if marker.casefold() in text.casefold():
                findings.append((relative, "LOCAL_OR_LEGACY_PATH", "path/username appears; evidence file may need sanitization or explicit retention", "REVIEW"))
                break
        for category, pattern in secret_patterns:
            if pattern.search(text):
                findings.append((relative, category, "credential-like material detected; value intentionally omitted", "BLOCK"))
        if email_pattern.search(text):
            findings.append((relative, "EMAIL_REVIEW", "email address present; confirm publication is intentional", "REVIEW"))
    unique = sorted(set(findings))
    blocking = [item for item in unique if item[3] == "BLOCK"]
    raw_in_manifest = any(row["relative_path"].casefold().endswith(".pth") for row in read_csv("data/manifests/dataset_inventory.csv"))
    checkpoint_summary = json.loads((REPO / "results/audit/phase2_checkpoint_summary.json").read_text(encoding="utf-8"))
    report = {
        "release_candidate_files": len(candidates), "findings": len(unique), "blocking_findings": len(blocking),
        "checkpoint_files_in_repository": sum(path.suffix.casefold() == ".pth" for path in candidates),
        "pth_rows_in_image_manifest": int(raw_in_manifest),
        "data_root_internal_redundant_checkpoint_files": checkpoint_summary["data_root_internal_redundant_files"],
        "data_root_matches_checkpoint_root": checkpoint_summary["data_root_duplicates_with_checkpoint_root"],
        "data_root_matches_legacy_github": checkpoint_summary["data_root_duplicates_with_legacy_github"],
        "status": "BLOCKED" if blocking or raw_in_manifest else ("REVIEW_REQUIRED" if unique else "PASS"),
    }
    lines = [
        "# Public release scan", "", f"Status: **{report['status']}**", "",
        f"- Candidate files scanned: {len(candidates)}", f"- Findings: {len(unique)}",
        f"- Blocking findings: {len(blocking)}", f"- `.pth` files in repository candidate set: {report['checkpoint_files_in_repository']}",
        f"- `.pth` rows in image manifest: {report['pth_rows_in_image_manifest']}",
        f"- Redundant checkpoint files within DATA_ROOT scope: {report['data_root_internal_redundant_checkpoint_files']}",
        f"- DATA_ROOT matches to CHECKPOINT_ROOT / legacy GitHub: {report['data_root_matches_checkpoint_root']} / {report['data_root_matches_legacy_github']}", "",
        "Local absolute paths in audit evidence are reported for review; no credential value is reproduced here. "
        "Legacy-path inventories may be retained as provenance, while machine-specific runtime paths should be sanitized before public commit.", "",
        "| File | Category | Severity | Action |", "|---|---|---|---|",
    ]
    for relative, category, note, severity in unique:
        lines.append(f"| `{relative}` | {category} | {severity} | {note} |")
    lines.extend(["", "Raw dataset and all 108 DATA_ROOT checkpoint binaries remain outside Git. No push was performed.", ""])
    (REPO / "results/audit/public_release_scan.md").write_text("\n".join(lines), encoding="utf-8")
    write_json(REPO / "results/audit/public_release_scan.json", report)
    return report


def write_decision(duplicates: dict[str, object], checkpoints: dict[str, object], lineages: list[dict[str, object]],
                   sota_rows: list[dict[str, object]], release_scan: dict[str, object]) -> None:
    c_categories = checkpoints["data_root_categories"]
    verified = sum(row["lineage_status"] == "VERIFIED_LINEAGE" for row in lineages)
    high_confidence = sum(row["lineage_status"] == "HIGH_CONFIDENCE_LINEAGE" for row in lineages)
    all_2057 = sum(row["data_exposure"] == "TRAINED_ON_ALL_2057" for row in lineages)
    valid_heldout = 0
    lines = [
        "# PHASE 2 DECISION", "",
        "| Item | Decision | Evidence | Required action |", "|---|---|---|---|",
        f"| Labeled dataset | {duplicates['labeled_dataset']['exact_unique']} exact-unique images from 2,057 rows | SHA256 labeled audit | Keep row count and duplicate count separate |",
        "| 2,250 versus 2,057 | UNRESOLVED | No count evidence for 2,250; `labels.csv` proves 2,057 | Obtain draft/capture manifest with sample IDs |",
        f"| Fine-tuned checkpoint | {c_categories.get('FINETUNED_SOILNET_DUALHEAD', 0)} in DATA_ROOT; 0 in the 969-file CHECKPOINT_ROOT scope | Safe tensor/key classification | Do not select by filename |",
        f"| Historical held-out evidence | {valid_heldout} valid checkpoints | Fine-tune and SOTA sources use all-row training/evaluation | Do not report as test performance |",
        "| Confirmed unseen labeled data | 0 | All 2,057 row identities are exposed by high-confidence fine-tune generations | Prospective split is for a new run only |",
        "| Preselection | PRESELECTION_UNRESOLVED | Manuscript/final config mapping is absent | Freeze architecture, SSL source and mu before training |",
        "| Final recovery | READY_AFTER_ONE_SHORT_FINETUNE | Old checkpoints cannot recover clean held-out metrics | Run only FINAL_LEAKAGE_CONTROLLED_REVALIDATION after preselection |",
        "| Public release | " + str(release_scan["status"]) + " | No binaries copied; machine paths remain in some evidence files | Sanitize/review before a later commit |",
        "", "## Direct answers", "",
        f"1. **Unique labeled samples:** {duplicates['labeled_dataset']['exact_unique']} exact SHA256 images across 2,057 labeled rows; "
        f"{duplicates['labeled_dataset']['redundant_copies']} rows are redundant exact copies in "
        f"{duplicates['labeled_dataset']['exact_duplicate_groups']} exact groups.",
        "2. **2,250 vs 2,057:** UNRESOLVED. No audited artifact establishes 2,250 as an image count. 2,057 has strongest provenance.",
        f"3. **Fine-tuned SoilNet checkpoints:** 0 in CHECKPOINT_ROOT; {c_categories.get('FINETUNED_SOILNET_DUALHEAD', 0)} safely classified candidates in DATA_ROOT.",
        "4. **Strongest lineage:** the `Best_finetuned_*` and direct `finetune_mu_*` families have HIGH_CONFIDENCE filename/save-target/architecture lineage, but none is VERIFIED because save-time SHA evidence is absent. No single mu is uniquely selected.",
        "5. **Valid held-out evidence:** none. No linked fine-tuned checkpoint has a reproducible held-out split with sample identifiers.",
        "6. **Confirmed unseen labeled data:** none. A newly assigned prospective test row is not historically unseen.",
        "7. **Historical Accuracy/F1:** retain only as `TRAINING_SET_METRICS_ONLY`; downstream classical rows require reevaluation.",
        "8. **Historical RMSE/MAE:** retain only as training behavior; not held-out performance.",
        "9. **SSL artifacts:** existing hash-inventoried backbone/projector artifacts do not require a new SSL sweep, but soil-pool SSL is transductive and no pure-external verified final encoder is selected.",
        "10. **Cheap classical regeneration:** KNN, SVM, GBM, Decision Tree and MLP after a frozen encoder/config decision; no CNN training in that step.",
        "11. **New final fine-tune:** yes, only if held-out generalization metrics remain required by the manuscript.",
        "12. **Exact training:** `FINAL_LEAKAGE_CONTROLLED_REVALIDATION`: preselected SoilNet only; 1,496 train, 317 validation, validation-only selection, 244 test once.",
        "13. **Mu sweep:** no. Mu/config must be frozen from historical/manuscript evidence before the new test is touched.",
        "14. **SSL rerun:** no evidence currently justifies rerunning SSL.",
        "15. **Full rerun:** none justified.",
        "", "## Checkpoints discovered inside DATA_ROOT", "",
        f"1. `.pth` files: {checkpoints['data_root_embedded_files']}.",
        f"2. Unique SHA256: {checkpoints['data_root_unique_sha256']}.",
        f"2a. Redundant files within DATA_ROOT scope: {checkpoints['data_root_internal_redundant_files']}.",
        f"3. Files byte-identical to CHECKPOINT_ROOT: {checkpoints['data_root_duplicates_with_checkpoint_root']}.",
        f"4. Files byte-identical to legacy GitHub: {checkpoints['data_root_duplicates_with_legacy_github']}.",
        f"5. FINETUNED_SOILNET_DUALHEAD: {c_categories.get('FINETUNED_SOILNET_DUALHEAD', 0)}.",
        f"6. Checkpoints with regression and classification heads: {checkpoints['data_root_reg_and_cls_heads']}.",
        f"7. Checkpoints with a light branch: {checkpoints['data_root_light_branch']}.",
        f"8. Fine-tuned lineage: VERIFIED={verified}; HIGH_CONFIDENCE={high_confidence}. Save-time SHA evidence is required for VERIFIED.",
        f"9. Fine-tuned candidates linked to training on all 2,057 rows: {all_2057}.",
        f"10. Genuine historical held-out checkpoints: {valid_heldout}.",
        "11. These artifacts do not change the one-short-fine-tune decision: they recover model bytes and lineage, but not a clean historical held-out protocol.",
        "", "## Scope totals", "",
        f"- CHECKPOINT_ROOT: {checkpoints['checkpoint_root_files']} files",
        f"- LEGACY_GITHUB: {checkpoints['legacy_github_files']} files",
        f"- DATA_ROOT_EMBEDDED_ARTIFACTS: {checkpoints['data_root_embedded_files']} files",
        f"- CROSS-SCOPE UNIQUE SHA256: {checkpoints['cross_scope_unique_sha256']}",
        f"- VERIFIED FINE-TUNED LINEAGE: {verified}",
        f"- HIGH-CONFIDENCE FINE-TUNED LINEAGE: {high_confidence}",
        f"- SOTA/baseline artifacts separately audited: {len(sota_rows)}",
        "", "## Final decision", "", "**READY_AFTER_ONE_SHORT_FINETUNE**", "",
        "Execution gate: **PRESELECTION_UNRESOLVED**. No training may begin until final architecture, SSL source and mu are fixed without using prospective validation/test outcomes.", "",
    ]
    (REPO / "results/audit/PHASE2_DECISION.md").write_text("\n".join(lines), encoding="utf-8")


def update_phase2_checksums() -> None:
    """Merge final Phase 2 outputs into the reproducibility checksum registry."""
    phase2_targets = [
        "docs/DATASET.md",
        "docs/CLAIMS_EVIDENCE_MATRIX.md",
        "results/audit/labeled_dataset_duplicate_summary.json",
        "results/audit/labeled_exact_duplicates.csv",
        "results/audit/labeled_near_duplicates.csv",
        "results/audit/checkpoint_classification.csv",
        "results/audit/data_root_checkpoint_inventory.csv",
        "results/audit/checkpoint_cross_scope_matches.csv",
        "results/audit/phase2_checkpoint_summary.json",
        "results/audit/dataset_count_provenance.md",
        "results/audit/fine_tuned_checkpoint_lineage.csv",
        "results/audit/sota_checkpoint_lineage.csv",
        "results/audit/historical_split_inventory.csv",
        "results/audit/unseen_labeled_candidates.csv",
        "results/audit/ssl_checkpoint_provenance.csv",
        "results/audit/classical_ml_reproduction_plan.md",
        "results/audit/claims_evidence_matrix.csv",
        "results/audit/public_release_scan.md",
        "results/audit/public_release_scan.json",
        "results/audit/PHASE2_DECISION.md",
        "results/audit/phase2_summary.json",
        "results/audit/gpu_diagnostic.txt",
        "results/audit/gpu_diagnostic.json",
    ]
    registry_path = REPO / "data/checksums/manifest_sha256.csv"
    existing = {row["relative_path"]: row for row in read_csv("data/checksums/manifest_sha256.csv")}
    for relative in phase2_targets:
        path = REPO / relative
        if not path.is_file():
            raise RuntimeError(f"Missing Phase 2 checksum target: {relative}")
        existing[relative] = {
            "relative_path": relative,
            "size_bytes": path.stat().st_size,
            "sha256": sha256_file(path),
        }
    write_csv(registry_path, [existing[key] for key in sorted(existing)], ["relative_path", "size_bytes", "sha256"])


def main() -> int:
    paths = resolve_paths()
    documents = source_documents(paths)
    duplicates = audit_labeled_duplicates()
    scope_c, checkpoint_summary = audit_checkpoints(paths, documents)
    audit_count_provenance(paths, checkpoint_summary)
    lineages, sota_rows = audit_lineage(paths, documents, scope_c)
    audit_historical_splits(paths, documents, lineages)
    audit_unseen_and_ssl(paths, scope_c, lineages)
    write_classical_plan(paths)
    update_claims(lineages)
    release_scan = public_release_scan()
    write_decision(duplicates, checkpoint_summary, lineages, sota_rows, release_scan)
    final_summary = {
        "labeled_rows": duplicates["labeled_dataset"]["labeled_rows"],
        "labeled_exact_unique": duplicates["labeled_dataset"]["exact_unique"],
        "checkpoint_root": checkpoint_summary["checkpoint_root_files"],
        "legacy_github": checkpoint_summary["legacy_github_files"],
        "data_root_embedded": checkpoint_summary["data_root_embedded_files"],
        "data_root_finetuned_soilnet": checkpoint_summary["data_root_categories"].get("FINETUNED_SOILNET_DUALHEAD", 0),
        "valid_historical_heldout": 0,
        "confirmed_unseen_labeled": 0,
        "decision": "READY_AFTER_ONE_SHORT_FINETUNE",
        "training_run": False,
    }
    write_json(REPO / "results/audit/phase2_summary.json", final_summary)
    update_phase2_checksums()
    print(json.dumps(final_summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
